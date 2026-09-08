#!/usr/bin/env python3
"""
Lists every IGHD gene that has more than one observed allele in the real
KIARVA genotype sample (the only genes where a geographic allele-frequency
comparison is even possible - a single-allele gene has no diversity to
test), alongside the geographic significance result already computed in
build_dregion_consensus.py.

Usage:
    python scripts/build_multi_allele_list.py [--cache-dir .cache] [--out results/Coklu_Alelli_Genler.xlsx]
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from vdj_bias.kiarva_genotypes import load_d_gene_rows

FONT = "Arial"

# significance results from build_dregion_consensus.py's geographic chi-square test
GEOGRAPHIC_SIGNIFICANCE = {
    "IGHD4-4": ("2.11e-03", True),
    "IGHD2-21": ("9.67e-23", True),
    "IGHD2-2": ("2.78e-42", True),
    "IGHD3-10": ("4.90e-03", True),
    "IGHD3-3": ("1.04e-34", True),
    "IGHD3-16": ("7.99e-03", True),
    "IGHD4-17": ("9.98e-01", False),
    "IGHD5-18": ("5.30e-01", False),
    "IGHD2-8": ("5.30e-01", False),
    "IGHD5-18/5-5": ("4.79e-01", False),
}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cache-dir", default=".cache")
    ap.add_argument("--out", default="results/Coklu_Alelli_Genler.xlsx")
    args = ap.parse_args()

    d_rows = load_d_gene_rows(Path(args.cache_dir) / "kiarva_genotypes")
    per_gene = d_rows.groupby("gene")["base_db_name"].unique()
    multi = {g: sorted(a) for g, a in per_gene.items() if len(a) >= 2}
    multi_sorted = sorted(multi.items(), key=lambda kv: -len(kv[1]))

    wb = Workbook()
    ws = wb.active
    ws.title = "Coklu_Alelli_Genler"

    ws.merge_cells("A1:E1")
    ws["A1"] = "Birden Fazla Aleli Olan D Genleri (Ornleklemde)"
    ws["A1"].font = Font(name=FONT, size=13, bold=True)

    headers = ["D Geni", "Alel Sayisi", "Alel Isimleri", "Cografi p (BH)", "Anlamli mi"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2E5B8A")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    r = 4
    for gene, alleles in multi_sorted:
        p_val, is_sig = GEOGRAPHIC_SIGNIFICANCE.get(gene, ("test edilmedi", None))
        ws.cell(row=r, column=1, value=gene).font = Font(name=FONT, bold=True)
        ws.cell(row=r, column=2, value=len(alleles)).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=3, value=", ".join(alleles))
        ws.cell(row=r, column=4, value=p_val).alignment = Alignment(horizontal="center")
        sig_cell = ws.cell(row=r, column=5, value="EVET *" if is_sig else "hayir" if is_sig is False else "-")
        sig_cell.alignment = Alignment(horizontal="center")
        if is_sig:
            sig_cell.font = Font(name=FONT, bold=True, color="1F6B2C")
        r += 1

    ws.cell(row=r + 1, column=1, value=f"Toplam: {len(multi_sorted)} / {len(per_gene)} gen birden fazla alele sahip").font = Font(
        name=FONT, italic=True, size=9
    )

    for i, w in enumerate([14, 12, 40, 16, 12], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"

    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"Yazildi: {out_path}")


if __name__ == "__main__":
    main()
