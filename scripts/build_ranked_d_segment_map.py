#!/usr/bin/env python3
"""
Ranks the 26 IGHD genes that have any real SARP data (IGHD4-11 has none on
either flank and is excluded, not guessed) by a combined "D segment VDJ
inclusion score" - the mean of its 5' (V-proximal) and 3' (J-proximal) mean
SARP scores (or just the one side available, for the 9 genes missing 5' data).
Rank 1 = highest combined score = most likely to end up in the final VDJ join.

Produces two deliverables sharing the SAME gene order (fixed by the
all-people ranking, so the two are directly comparable position-by-position):

  1. results/D_Segment_Siralama_TumInsanlar.xlsx
     One track, pooled across every real individual in the sample.

  2. results/D_Segment_Siralama_Cografi.xlsx
     Three tracks (Africa/Asia/Europe), each real individual's own combined
     score compared across groups with Kruskal-Wallis + BH-corrected
     pairwise Mann-Whitney (real 410/group cohorts) - '*' marks a gene where
     that test came back significant (p_adj<0.05).

Usage:
    python scripts/build_ranked_d_segment_map.py [--cache-dir .cache] [--out-dir results]
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from vdj_bias.analysis import run_statistics, score_cohorts
from vdj_bias.kiarva_genotypes import GROUP_TO_SUPERPOPS, build_group_cohorts, build_rss_reference_table, build_genotype_cohort, load_d_gene_rows
from vdj_bias.sarp_scores import load_sarp_scores
from vdj_bias.vdjbase_client import build_rss_reference_table as build_vdjbase_rss_table

FONT = "Arial"
ALPHA = 0.05
DARK = "1F3864"
LIGHT = "FFFFFF"


def load_shared_inputs(cache_dir: Path):
    d_rows = load_d_gene_rows(cache_dir / "kiarva_genotypes")
    sarp_scores = load_sarp_scores(cache_dir / "sarp")
    primary_rss = build_rss_reference_table(d_rows)
    vdjbase_rss = build_vdjbase_rss_table(cache_dir / "vdjbase")
    have = set(zip(primary_rss["gene"], primary_rss["allele"], primary_rss["side"]))
    extra = vdjbase_rss[~vdjbase_rss.apply(lambda r: (r["gene"], r["allele"], r["side"]) in have, axis=1)]
    rss_ref = pd.concat([primary_rss, extra], ignore_index=True)
    genes = sorted(d_rows["gene"].unique())
    return d_rows, sarp_scores, rss_ref, genes


def combined_per_person_scores(cohorts, rss_ref, sarp_scores, genes) -> pd.DataFrame:
    """One row per (group, gene, individual_idx): combined_score = mean of
    that person's 5'-side and 3'-side scores, or whichever one exists."""
    frames = []
    for side in ("5", "3"):
        scored = score_cohorts(cohorts, rss_ref, sarp_scores, side=side)
        if scored.empty:
            continue
        scored = scored.rename(columns={"sarp_score": f"score_{side}"})
        frames.append(scored.set_index(["group", "gene", "individual_idx"])[f"score_{side}"])
    if not frames:
        return pd.DataFrame(columns=["group", "gene", "individual_idx", "sarp_score"])
    merged = pd.concat(frames, axis=1)
    merged["sarp_score"] = merged.mean(axis=1, skipna=True)
    return merged.reset_index()[["group", "gene", "individual_idx", "sarp_score"]].dropna(subset=["sarp_score"])


def rank_genes(all_people_scored: pd.DataFrame) -> pd.DataFrame:
    means = all_people_scored.groupby("gene")["sarp_score"].agg(mean="mean", n="count").reset_index()
    means = means.sort_values("mean", ascending=False).reset_index(drop=True)
    means["rank"] = means.index + 1
    return means


def style_header(ws, cell_range, fill_color="2E5B8A"):
    for row in ws[cell_range]:
        for c in row:
            c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=fill_color)
            c.alignment = Alignment(horizontal="center")


def build_all_people_workbook(ranked: pd.DataFrame, out_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Siralama_TumInsanlar"

    ws.merge_cells("A1:F1")
    ws["A1"] = "D Segmentleri: VDJ'ye Katilma Skoruna Gore Siralama (1=en yuksek), Tum Orneklem"
    ws["A1"].font = Font(name=FONT, size=13, bold=True)
    ws.merge_cells("A2:F2")
    ws["A2"] = (
        "Skor = kisinin 5' (V-tarafi) ve 3' (J-tarafi) SARP skorlarinin ortalamasi (sadece bir taraf "
        "veriliyorsa o taraf). IGHD4-11 hicbir tarafta veri olmadigi icin sirlamaya alinmadi."
    )
    ws["A2"].font = Font(name=FONT, size=9, italic=True, color="595959")

    headers = ["Sira", "D Geni", "Konsensus Skor (tum orneklem ort.)", "Kisi Sayisi (N)"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=4, column=c, value=h)
    style_header(ws, "A4:D4")

    r = 5
    value_cells = []
    for _, row in ranked.iterrows():
        ws.cell(row=r, column=1, value=int(row["rank"])).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=2, value=row["gene"]).font = Font(name=FONT, bold=True)
        vcell = ws.cell(row=r, column=3, value=round(float(row["mean"]), 4))
        vcell.alignment = Alignment(horizontal="center")
        value_cells.append(f"C{r}")
        ws.cell(row=r, column=4, value=int(row["n"])).alignment = Alignment(horizontal="center")
        r += 1

    rule = ColorScaleRule(start_type="min", start_color=DARK, end_type="max", end_color=LIGHT)
    ws.conditional_formatting.add(" ".join(value_cells), rule)

    for i, w in enumerate([8, 14, 30, 14], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def build_geographic_workbook(ranked: pd.DataFrame, region_scored: pd.DataFrame, sig: dict[str, bool], out_path: Path):
    wide = region_scored.groupby(["gene", "group"])["sarp_score"].agg(["mean", "count"])

    wb = Workbook()
    ws = wb.active
    ws.title = "Siralama_Cografi"

    ws.merge_cells("A1:H1")
    ws["A1"] = "D Segmentleri: Ayni Siralama, Kitasal Gruplara Gore (Africa / Asia / Europe)"
    ws["A1"].font = Font(name=FONT, size=13, bold=True)
    ws.merge_cells("A2:H2")
    ws["A2"] = (
        "Sira, 'Tum Insanlar' dosyasindaki ile ayni (tum orneklemin konsensus skoruna gore sabitlendi), "
        "boylece gruplar arasi karsilastirma pozisyon pozisyon yapilabilir. '*' = Kruskal-Wallis + BH "
        "duzeltmesiyle 3 grup arasinda istatistiksel olarak anlamli fark (p<0.05), gercek 410 kisi/grup ile."
    )
    ws["A2"].font = Font(name=FONT, size=9, italic=True, color="595959")

    groups = list(GROUP_TO_SUPERPOPS.keys())
    headers = ["Sira", "D Geni"] + [f"{g} Ort." for g in groups] + [f"{g} N" for g in groups] + ["Anlamli mi"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=4, column=c, value=h)
    style_header(ws, f"A4:{get_column_letter(len(headers))}4")

    r = 5
    value_cells = []
    for _, row in ranked.iterrows():
        gene = row["gene"]
        ws.cell(row=r, column=1, value=int(row["rank"])).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=2, value=gene).font = Font(name=FONT, bold=True)
        is_sig = sig.get(gene, False)
        for i, g in enumerate(groups):
            col = 3 + i
            mean_val = wide["mean"].get((gene, g))
            cell = ws.cell(row=r, column=col)
            if pd.isna(mean_val):
                cell.value = "veri yok"
                cell.font = Font(name=FONT, size=9, italic=True, color="808080")
            else:
                cell.value = f"{round(float(mean_val), 4)}{' *' if is_sig else ''}"
                cell.font = Font(name=FONT, bold=is_sig)
                value_cells.append(f"{get_column_letter(col)}{r}")
            cell.alignment = Alignment(horizontal="center")
        for i, g in enumerate(groups):
            col = 3 + len(groups) + i
            n_val = wide["count"].get((gene, g))
            ws.cell(row=r, column=col, value=None if pd.isna(n_val) else int(n_val)).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=3 + 2 * len(groups), value="EVET" if is_sig else "hayir").alignment = Alignment(horizontal="center")
        r += 1

    rule = ColorScaleRule(start_type="min", start_color=DARK, end_type="max", end_color=LIGHT)
    ws.conditional_formatting.add(" ".join(value_cells), rule)

    for i, w in enumerate([8, 14, 12, 12, 12, 10, 10, 10, 12], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"

    n_sig = sum(sig.values())
    ws.cell(row=r + 1, column=1, value=f"Anlamli gen sayisi: {n_sig} / {len(sig)} test edilen").font = Font(
        name=FONT, size=10, italic=True
    )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cache-dir", default=".cache")
    ap.add_argument("--out-dir", default="results")
    ap.add_argument("--n-per-group", type=int, default=410)
    ap.add_argument("--seed", type=int, default=42)
    args = ap.parse_args()

    cache_dir = Path(args.cache_dir)
    out_dir = Path(args.out_dir)

    print("[1/4] Ortak veriler yukleniyor (SARP skorlari, genotip, RSS referans tablosu)...")
    d_rows, sarp_scores, rss_ref, genes = load_shared_inputs(cache_dir)

    print("[2/4] Tum gercek bireyler icin birlesik (5'+3' ortalama) skor hesaplaniyor...")
    all_cases = d_rows["case"].unique().tolist()
    all_cohort = {"All": build_genotype_cohort(d_rows, all_cases, genes)}
    all_scored = combined_per_person_scores(all_cohort, rss_ref, sarp_scores, genes)
    ranked = rank_genes(all_scored)
    print(f"      {len(ranked)} gen siralandi (veri olmayanlar haric).")

    print("[3/4] Cografi kohortlar (gercek 410/grup) icin ayni birlesik skor + istatistik...")
    geo_cohorts = build_group_cohorts(d_rows, genes, n_per_group=args.n_per_group, seed=args.seed)
    region_scored = combined_per_person_scores(geo_cohorts, rss_ref, sarp_scores, genes)
    kw_df, _ = run_statistics(region_scored, alpha=ALPHA)
    sig = {row["gene"]: bool(row["significant"]) and pd.notna(row["p_value"]) for _, row in kw_df.iterrows()}
    n_sig = sum(sig.values())
    print(f"      {n_sig} / {len(sig)} gen istatistiksel olarak anlamli (p<0.05, BH duzeltmeli).")

    print("[4/4] Excel dosyalari yaziliyor...")
    build_all_people_workbook(ranked, out_dir / "D_Segment_Siralama_TumInsanlar.xlsx")
    build_geographic_workbook(ranked, region_scored, sig, out_dir / "D_Segment_Siralama_Cografi.xlsx")
    print(f"Yazildi: {out_dir}/D_Segment_Siralama_TumInsanlar.xlsx ve {out_dir}/D_Segment_Siralama_Cografi.xlsx")


if __name__ == "__main__":
    main()
