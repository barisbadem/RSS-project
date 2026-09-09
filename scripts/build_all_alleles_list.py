#!/usr/bin/env python3
"""
Full list of every real allele observed in the KIARVA genotype sample,
across all 28 IGHD genes, with its D-REGION sequence, how many real
short-read / long-read observations support it, and - for each of the 3
possible reading frames (RF1/RF2/RF3, i.e. starting the codon split at
position 0, 1 or 2 of the sequence) - the translated amino-acid string,
with any stop codon shown as "*" and flagged. A D gene's sequence itself
never shifts frame; which frame a real B cell actually uses at the VDJ
junction depends on how much exonuclease trimming/N-addition happens
there (see prior discussion) - this table just shows what each of the 3
theoretically possible frames would produce for every real allele.

Sequence lookup is keyed by (gene, base_allele) rather than the allele's
own db_name, because some near-identical duplicated genes (e.g. IGHD4-4 /
IGHD4-11) only get resolved to a clean single-gene name in the longer
flank-extended reads - their plain short reads carry an ambiguous
compound name like "IGHD4-11*01/IGHD4-4*01". A handful of rare alleles
that were only ever seen in a flank-extended read (never their own short
call) are marked as such rather than guessing their exact core boundary.

Usage:
    python scripts/build_all_alleles_list.py [--cache-dir .cache] [--out results/Tum_Alleller_Listesi.xlsx]
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from vdj_bias.kiarva_genotypes import load_d_gene_rows

FONT = "Arial"
NOT_FOUND_LABEL = "(sadece uzun okumada mevcut, cekirdek ayri cikarilamadi)"

# standard genetic code (DNA codon -> single-letter amino acid, "*" = stop)
CODON_TABLE = {
    "TTT": "F", "TTC": "F", "TTA": "L", "TTG": "L",
    "CTT": "L", "CTC": "L", "CTA": "L", "CTG": "L",
    "ATT": "I", "ATC": "I", "ATA": "I", "ATG": "M",
    "GTT": "V", "GTC": "V", "GTA": "V", "GTG": "V",
    "TCT": "S", "TCC": "S", "TCA": "S", "TCG": "S",
    "CCT": "P", "CCC": "P", "CCA": "P", "CCG": "P",
    "ACT": "T", "ACC": "T", "ACA": "T", "ACG": "T",
    "GCT": "A", "GCC": "A", "GCA": "A", "GCG": "A",
    "TAT": "Y", "TAC": "Y", "TAA": "*", "TAG": "*",
    "CAT": "H", "CAC": "H", "CAA": "Q", "CAG": "Q",
    "AAT": "N", "AAC": "N", "AAA": "K", "AAG": "K",
    "GAT": "D", "GAC": "D", "GAA": "E", "GAG": "E",
    "TGT": "C", "TGC": "C", "TGA": "*", "TGG": "W",
    "CGT": "R", "CGC": "R", "CGA": "R", "CGG": "R",
    "AGT": "S", "AGC": "S", "AGA": "R", "AGG": "R",
    "GGT": "G", "GGC": "G", "GGA": "G", "GGG": "G",
}


def translate_frame(seq: str, frame: int) -> str:
    """Translate seq starting at 0-based offset `frame` (0, 1 or 2). Any
    leftover 1-2 bases that don't complete a final codon are dropped."""
    seq = seq.upper()
    aas = []
    for i in range(frame, len(seq) - 2, 3):
        codon = seq[i : i + 3]
        aas.append(CODON_TABLE.get(codon, "X"))
    return "".join(aas)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cache-dir", default=".cache")
    ap.add_argument("--out", default="results/Tum_Alleller_Listesi.xlsx")
    args = ap.parse_args()

    d_rows = load_d_gene_rows(Path(args.cache_dir) / "kiarva_genotypes")
    short_rows = d_rows[~d_rows["is_long"]]

    core_seq_by_gene_allele: dict[tuple[str, str], str] = {}
    for gene, base_allele, seq in zip(short_rows["gene"], short_rows["base_allele"], short_rows["sequence"]):
        core_seq_by_gene_allele.setdefault((gene, base_allele), seq)

    combos = d_rows[["gene", "base_allele", "base_db_name"]].drop_duplicates()
    counts_short = short_rows.groupby("base_db_name").size()
    counts_long = d_rows[d_rows["is_long"]].groupby("base_db_name").size()

    rows = []
    for _, row in combos.iterrows():
        gene, base_allele, base_db_name = row["gene"], row["base_allele"], row["base_db_name"]
        seq = core_seq_by_gene_allele.get((gene, base_allele), NOT_FOUND_LABEL)
        rows.append(
            (gene, base_db_name, seq, int(counts_short.get(base_db_name, 0)), int(counts_long.get(base_db_name, 0)))
        )
    rows = sorted(set(rows), key=lambda r: (r[0], r[1]))

    wb = Workbook()
    ws = wb.active
    ws.title = "Tum_Alleller"
    ws.merge_cells("A1:H1")
    ws["A1"] = f"Orneklemdeki TUM Gercek Aleller ({len(rows)} satir) - 3 Olasi Okuma Cercevesinde Aminoasit Dizisi"
    ws["A1"].font = Font(name=FONT, size=13, bold=True)
    ws.merge_cells("A2:H2")
    ws["A2"] = (
        "RF1/RF2/RF3: dizinin 1., 2. ve 3. bazindan baslayarak kodon okuma. Gercek bir B hucresinde hangi "
        "cercevenin kullanildigi VDJ birlesme noktasindaki kirpilma/ekleme miktarina bagli (bkz. onceki aciklama); "
        "burada sadece 3 cercevenin de teorik olarak neye cevrildigi gosteriliyor. '*' = dur kodonu."
    )
    ws["A2"].font = Font(name=FONT, size=9, italic=True, color="595959")

    headers = [
        "D Geni",
        "Alel Adi",
        "D-REGION Dizisi",
        "Kisa-Okuma Gozlemi",
        "Uzun-Okuma Gozlemi",
        "RF1 (Aminoasit)",
        "RF2 (Aminoasit)",
        "RF3 (Aminoasit)",
    ]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2E5B8A")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    r = 5
    for gene, allele, seq, n_short, n_long in rows:
        has_seq = seq != NOT_FOUND_LABEL
        frames = [translate_frame(seq, f) for f in range(3)] if has_seq else ["-", "-", "-"]
        vals = [gene, allele, seq, n_short, n_long, frames[0], frames[1], frames[2]]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = Alignment(horizontal="center" if c not in (3, 6, 7, 8) else "left")
            if c == 1:
                cell.font = Font(name=FONT, bold=True)
            if c == 3 and v == NOT_FOUND_LABEL:
                cell.font = Font(name=FONT, italic=True, color="808080")
            if c in (6, 7, 8):
                if has_seq and "*" in v:
                    cell.font = Font(name=FONT, color="C0392B")
                else:
                    cell.font = Font(name=FONT)
        r += 1

    ws.cell(
        row=r + 1, column=1, value=f"Toplam: {len(rows)} satir, {len(set(x[0] for x in rows))} gen"
    ).font = Font(name=FONT, italic=True, size=9)

    for i, w in enumerate([14, 24, 42, 16, 16, 26, 26, 26], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"

    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"Yazildi: {out_path}")


if __name__ == "__main__":
    main()
