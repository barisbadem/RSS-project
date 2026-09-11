#!/usr/bin/env python3
"""
For every real D gene and every real allele of that gene, stacked
vertically: the allele's D-REGION DNA sequence, then its translation in
each of the 3 possible reading frames (RF1/RF2/RF3), one frame per row -
then the next allele of the same gene, then the next gene, and so on for
every gene and allele in the sample (not just the geographically
significant ones).

Usage:
    python scripts/build_allele_frames_stacked.py [--cache-dir .cache] [--out results/Alel_Cerceve_Dizileri.xlsx]
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from scripts.build_all_alleles_list import NOT_FOUND_LABEL, translate_frame
from vdj_bias.kiarva_genotypes import build_core_sequence_lookup, load_d_gene_rows

FONT = "Arial"


def _gene_sort_key(gene: str) -> tuple[int, str]:
    """Sort by real genomic position number where parseable (e.g. 'IGHD5-18/5-5'
    sorts as position 18), falling back to alphabetical for anything odd."""
    m = re.match(r"IGHD\d+-(\d+)", gene)
    return (int(m.group(1)) if m else 999, gene)


def gather_gene_alleles(d_rows) -> dict[str, list[tuple[str, str, int]]]:
    """{gene: [(allele_name, dna_sequence, n_people), ...]} most-common allele
    first. n_people counts real distinct (case) presence, short OR long read."""
    core_seq = build_core_sequence_lookup(d_rows)
    combos = d_rows[["gene", "base_allele", "base_db_name"]].drop_duplicates()
    presence = d_rows[["case", "gene", "base_db_name"]].drop_duplicates()
    counts = presence.groupby(["gene", "base_db_name"]).size()

    gene_alleles: dict[str, list[tuple[str, str, int]]] = {}
    for gene, base_allele, base_db_name in combos.itertuples(index=False):
        seq = core_seq.get((gene, base_allele), NOT_FOUND_LABEL)
        n = int(counts.get((gene, base_db_name), 0))
        gene_alleles.setdefault(gene, []).append((base_db_name, seq, n))

    for gene in gene_alleles:
        gene_alleles[gene].sort(key=lambda x: -x[2])
    return gene_alleles


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cache-dir", default=".cache")
    ap.add_argument("--out", default="results/Alel_Cerceve_Dizileri.xlsx")
    args = ap.parse_args()

    cache_dir = Path(args.cache_dir)
    print("[1/2] Veriler yukleniyor, gen basina alel listesi cikariliyor...")
    d_rows = load_d_gene_rows(cache_dir / "kiarva_genotypes")
    gene_alleles = gather_gene_alleles(d_rows)
    genes_sorted = sorted(gene_alleles, key=_gene_sort_key)

    print("[2/2] Excel yaziliyor...")
    wb = Workbook()
    ws = wb.active
    ws.title = "Alel_Cerceveleri"

    ws.merge_cells("A1:D1")
    ws["A1"] = "Her D Geni ve Her Alel Icin DNA Dizisi + 3 Okuma Cercevesinde Amino Asit Cevirisi"
    ws["A1"].font = Font(name=FONT, size=13, bold=True)
    ws.merge_cells("A2:D2")
    ws["A2"] = (
        "Her alel bloğu: DNA dizisi, sonra RF1/RF2/RF3 (1., 2. ve 3. bazdan baslayan kodon okumasi) "
        "amino asit cevirisi, alt alta. '*' = dur kodonu (kirmizi). n = o aleli tasidigi gozlenen gercek kisi sayisi."
    )
    ws["A2"].font = Font(name=FONT, size=9, italic=True, color="595959")

    gene_fill = PatternFill("solid", fgColor="2E5B8A")
    allele_fill = PatternFill("solid", fgColor="D9E8F5")
    label_font = Font(name=FONT, size=9, bold=True, color="595959")

    r = 4
    for gene in genes_sorted:
        gene_start_row = r
        for allele_name, seq, n in gene_alleles[gene]:
            has_seq = seq != NOT_FOUND_LABEL
            frames = [translate_frame(seq, f) for f in range(3)] if has_seq else ["-", "-", "-"]

            allele_start_row = r
            rows = [
                ("DNA", seq),
                ("RF1", frames[0]),
                ("RF2", frames[1]),
                ("RF3", frames[2]),
            ]
            for label, value in rows:
                ws.cell(row=r, column=3, value=label).font = label_font
                cell = ws.cell(row=r, column=4, value=value)
                cell.font = Font(name=FONT, size=10)
                cell.alignment = Alignment(horizontal="left")
                if label != "DNA":
                    if has_seq and "*" in value:
                        cell.font = Font(name=FONT, size=10, color="C0392B")
                elif not has_seq:
                    cell.font = Font(name=FONT, size=10, italic=True, color="808080")
                r += 1

            ws.merge_cells(start_row=allele_start_row, start_column=2, end_row=r - 1, end_column=2)
            acell = ws.cell(row=allele_start_row, column=2, value=f"{allele_name} (n={n})")
            acell.font = Font(name=FONT, size=10, bold=True)
            acell.fill = allele_fill
            acell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws.merge_cells(start_row=gene_start_row, start_column=1, end_row=r - 1, end_column=1)
        gcell = ws.cell(row=gene_start_row, column=1, value=gene)
        gcell.font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
        gcell.fill = gene_fill
        gcell.alignment = Alignment(horizontal="center", vertical="center")

        r += 1  # blank spacer row between genes

    for i, w in enumerate([14, 26, 6, 60], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"

    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"Yazildi: {out_path}")


if __name__ == "__main__":
    main()
