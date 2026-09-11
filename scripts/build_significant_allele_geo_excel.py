#!/usr/bin/env python3
"""
Excel version of the "D Alel Cografyasi" chart (published artifact): for the
7 D genes whose D-REGION allele composition is statistically significantly
different by geography (chi-square test of independence, BH-corrected
p<0.05 across all multi-allele genes), one row per allele with its real
observed frequency in each of Africa/Europe/Asia (~410-person-per-group
cohorts, haplotype-based) - same numbers, same gene/allele order as the
chart.

Usage:
    python scripts/build_significant_allele_geo_excel.py [--cache-dir .cache] [--out results/Anlamli_Allel_Cografyasi.xlsx]
"""
from __future__ import annotations

import argparse
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from scipy import stats

from scripts.build_dregion_consensus import haplotype_alleles_per_gene
from vdj_bias.analysis import _benjamini_hochberg
from vdj_bias.kiarva_genotypes import GROUP_TO_SUPERPOPS, build_group_cohorts, load_d_gene_rows

FONT = "Arial"
ALPHA = 0.05
GROUPS = list(GROUP_TO_SUPERPOPS.keys())  # ["Africa", "Europe", "Asia"]
GROUP_TR = {"Africa": "Afrika", "Europe": "Avrupa", "Asia": "Asya"}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cache-dir", default=".cache")
    ap.add_argument("--out", default="results/Anlamli_Allel_Cografyasi.xlsx")
    ap.add_argument("--n-per-group", type=int, default=410)
    ap.add_argument("--seed", type=int, default=42)
    args = ap.parse_args()

    cache_dir = Path(args.cache_dir)
    print("[1/2] Cografi kohortlar ve ki-kare testleri hesaplaniyor...")
    d_rows = load_d_gene_rows(cache_dir / "kiarva_genotypes")
    genes = sorted(d_rows["gene"].unique())
    geo_cohorts = build_group_cohorts(d_rows, genes, n_per_group=args.n_per_group, seed=args.seed)
    alleles_by_group = {g: haplotype_alleles_per_gene(geo_cohorts[g], genes) for g in GROUPS}

    gene_results = {}
    for gene in genes:
        allele_names = set()
        counts_by_group = {}
        for g in GROUPS:
            c = Counter(alleles_by_group[g][gene])
            counts_by_group[g] = c
            allele_names |= set(c.keys())
        allele_names = sorted(allele_names)
        if len(allele_names) < 2:
            continue
        table = np.array([[counts_by_group[g].get(a, 0) for a in allele_names] for g in GROUPS])
        if table.sum() == 0 or (table.sum(axis=0) == 0).any():
            continue
        try:
            _, p, _, _ = stats.chi2_contingency(table)
        except ValueError:
            continue
        gene_results[gene] = (p, allele_names, counts_by_group)

    tested = list(gene_results.keys())
    p_adj = _benjamini_hochberg(np.array([gene_results[g][0] for g in tested])) if tested else []
    p_adj_map = dict(zip(tested, p_adj))

    sig_genes = [g for g in tested if p_adj_map[g] < ALPHA]
    sig_genes.sort(key=lambda g: p_adj_map[g])

    print(f"      {len(sig_genes)} anlamli gen bulundu: {', '.join(sig_genes)}")

    print("[2/2] Excel yaziliyor...")
    wb = Workbook()
    ws = wb.active
    ws.title = "Anlamli_Allel_Cografyasi"

    ws.merge_cells("A1:F1")
    ws["A1"] = "Cografyaya Gore Istatistiksel Olarak Anlamli 7 D Geninin Alel Sikliklari"
    ws["A1"].font = Font(name=FONT, size=13, bold=True)
    ws.merge_cells("A2:F2")
    ws["A2"] = (
        "Gercek ~410 kisi/grup (Afrika/Avrupa/Asya) kohortlarindaki her alelin gercek gozlenen sikligi (%). "
        "Bu 7 gen, ki-kare bagimsizlik testinde (BH-duzeltmeli p<0.05) allel bilesiminin cografyaya gore "
        "anlamli sekilde farklilastigi genler - yayinlanan 'D Alel Cografyasi' grafik sayfasiyla ayni veri/siralama."
    )
    ws["A2"].font = Font(name=FONT, size=9, italic=True, color="595959")

    headers = ["D Geni", "Gen p (BH)", "Alel", "Afrika %", "Avrupa %", "Asya %"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2E5B8A")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    r = 5
    for gene in sig_genes:
        p_adj_val, allele_names, counts_by_group = gene_results[gene][0], gene_results[gene][1], gene_results[gene][2]
        # keep the exact same p-value shown as in the artifact/chart (BH-adjusted)
        p_adj_val = p_adj_map[gene]
        totals = {g: sum(counts_by_group[g].values()) for g in GROUPS}

        rows_for_gene = []
        for a in allele_names:
            pct = {g: (counts_by_group[g].get(a, 0) / totals[g] * 100 if totals[g] else 0) for g in GROUPS}
            rows_for_gene.append((a, pct))
        rows_for_gene.sort(key=lambda x: -max(x[1].values()))

        gene_start_row = r
        for allele, pct in rows_for_gene:
            vals = [
                gene,
                f"{p_adj_val:.2e}",
                allele,
                f"{pct['Africa']:.1f}%",
                f"{pct['Europe']:.1f}%",
                f"{pct['Asia']:.1f}%",
            ]
            for c, v in enumerate(vals, start=1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.alignment = Alignment(horizontal="center")
                if c == 1:
                    cell.font = Font(name=FONT, bold=True)
            r += 1
        if len(rows_for_gene) > 1:
            ws.merge_cells(start_row=gene_start_row, start_column=1, end_row=r - 1, end_column=1)
            ws.merge_cells(start_row=gene_start_row, start_column=2, end_row=r - 1, end_column=2)

    ws.cell(row=r + 1, column=1, value=f"Toplam {len(sig_genes)} anlamli gen, {r - 5} alel satiri.").font = Font(
        name=FONT, size=10, italic=True
    )

    for i, w in enumerate([14, 14, 24, 10, 10, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"

    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"Yazildi: {out_path}")


if __name__ == "__main__":
    main()
